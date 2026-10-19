"""Gera a documentação técnica; não participa da execução do dashboard.

Usa o ambiente de documentos da ferramenta, separado das dependências da aplicação.
"""
from pathlib import Path
from math import atan2, cos, sin
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'Arquitetura_do_dashboard_Heliopolis.docx'
QA = ROOT / 'verificacao'
QA.mkdir(exist_ok=True)

# O fluxograma é desenhado em alta resolução para permanecer legível no Word.
im = Image.new('RGB', (1800, 1840), 'white')
draw = ImageDraw.Draw(im)
font_path = 'C:/Windows/Fonts/arial.ttf'
font = ImageFont.truetype(font_path, 31)
small = ImageFont.truetype(font_path, 26)
bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 33)
def box(rect, title, body, fill='#EDF3F7'):
    draw.rounded_rectangle(rect, radius=15, fill=fill, outline='#547084', width=3)
    x0,y0,x1,y1=rect
    lines=[(title,bold)]+[(line,font) for line in body.split('\n')]
    y=(y0+y1)/2-((len(lines)-1)*40+34)/2
    for line,f in lines:
        draw.text(((x0+x1)/2,y),line,font=f,fill='#142A39',anchor='mt')
        y+=40
def arrow(points, label=None, labelpos=None):
    draw.line(points,fill='#435C6B',width=5)
    x,y=points[-1]; px,py=points[-2]; a=atan2(y-py,x-px)
    draw.polygon([(x,y),(x-20*cos(a-.45),y-20*sin(a-.45)),(x-20*cos(a+.45),y-20*sin(a+.45))],fill='#435C6B')
    if label: draw.text(labelpos,label,font=small,fill='#243E50',anchor='mm')

box((80,40,780,185),'Entrada da série ERA5','JSON salvo ou API Open-Meteo\nAPI somente se o arquivo não existir')
box((1020,40,1720,185),'Entrada das estações INMET','CSVs originais baixados pelo usuário\nPastas anuais e subpastas')
box((80,260,780,440),'dados.py','Campos e unidades • horários UTC\nGrade completa • valores inválidos\nLacunas preservadas')
box((1020,260,1720,440),'inmet_local.py e dados.py','Metadados • raio de 30 km\nLeitura dos CSVs • união por estação\nDuplicatas • cobertura e auditoria')
arrow([(430,185),(430,260)])
arrow([(1370,185),(1370,260)])
box((80,520,780,665),'painel.py','Recorte de 2021 a 2025\nFiltro de ano e limiar de temperatura')
box((1020,520,1720,665),'Séries independentes','Mirante e Interlagos\nMesmo período local do estudo')
arrow([(430,440),(430,520)])
arrow([(1370,440),(1370,520)])
box((80,760,780,945),'analise.py','Indicadores diários e noturnos\nCobertura mínima de 90%\nSequências de pelo menos três dias')
box((1020,760,1720,945),'contexto_regional.py','Recorta e chama indicadores_diarios\npara ERA5 e cada estação\nSeleciona as datas válidas em comum')
arrow([(430,665),(430,760)])
arrow([(1370,665),(1370,760)])
arrow([(780,590),(900,590),(900,720),(1190,720),(1190,760)],'ERA5 recortado',(1030,692))
box((80,1040,780,1220),'Resultados do contexto regional','Intensidade e frequência do calor\nPerfil horário • noites • umidade\nEpisódios e cartões')
box((1020,1040,1720,1220),'Comparação entre referências','Percentual de dias com calor\nMesmo conjunto de datas válidas\nSem fusão ou correção das séries')
arrow([(430,945),(430,1040)])
arrow([(1370,945),(1370,1040)])
box((350,1350,1450,1530),'Streamlit em uma única página','Gráficos e análises exploratórias\nAlterar um filtro recalcula resultados\nDados em sessão e cache evitam repetir a coleta',fill='#E4EEE8')
arrow([(430,1220),(430,1280),(690,1280),(690,1350)])
arrow([(1370,1220),(1370,1280),(1110,1280),(1110,1350)])
draw.text((900,1620),'Fora do fluxo atual',font=bold,fill='black',anchor='mt')
draw.text((900,1680),'Medições do Observatório e desfechos de saúde ainda não integrados.',font=font,fill='black',anchor='mt')
draw.text((900,1730),'As séries regionais não são medições dentro das moradias.',font=font,fill='black',anchor='mt')
flow = QA / 'fluxograma.png'; im.save(flow)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
sec.left_margin=Cm(2); sec.right_margin=Cm(2)
sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
for name,size in [('Normal',10.5),('Title',25),('Heading 1',17),('Heading 2',12),('Caption',9)]:
    style=doc.styles[name]; style.font.name='Calibri'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.12
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Normal'].paragraph_format.widow_control=True
# O modelo padrão pode herdar uma borda azul no estilo Title.
# Removemos bordas de parágrafo dos estilos para manter o título limpo.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
head=sec.header.paragraphs[0]; head.text='HELIÓPOLIS  |  DOCUMENTAÇÃO DA ARQUITETURA'
head.runs[0].font.size=Pt(8)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Versão de 18 outubro 2026  •  Página ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
for r in foot.runs: r.font.size=Pt(8)
doc.core_properties.title='Arquitetura e metodologia do dashboard de Heliópolis'
doc.core_properties.subject='Fluxo dos dados, ERA5, cálculos e responsabilidades dos módulos'
doc.core_properties.author='Projeto Heliópolis'
def p(t): return doc.add_paragraph(t)
def h(t): return doc.add_heading(t,level=2)
def page(t): doc.add_page_break(); doc.add_heading(t,level=1)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'24475C' if ri==0 else ('F0F4F7' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1.05
                for r in para.runs:
                    r.font.size=Pt(9.5)
                    if ri==0: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    p('')
    return t
def formula(text):
    # OMML mantém as expressões editáveis no Word, sem depender de imagens.
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); txt=OxmlElement('m:t'); txt.text=text
    run.append(txt); math.append(run); para._p.append(math)

doc.add_paragraph('Arquitetura e metodologia do dashboard de Heliópolis','Title')
p('Guia de estudo do fluxo de dados e dos cálculos implementados')
p('Este documento explica a implementação atual do dashboard: como os dados são obtidos, tratados e utilizados; quais cálculos produzem cada gráfico; e como os arquivos Python se comunicam. O período analisado é 2021–2025, definido como os cinco anos completos anteriores a 2026.')
h('Objetivo e alcance atual')
p('O objetivo do projeto é utilizar Ciência de Dados para analisar os impactos do estresse térmico na saúde e na qualidade de vida dos moradores de Heliópolis, integrando as medições do Observatório De Olho na Quebrada a fontes públicas externas.')
p('Como a base do Observatório não está disponível e não há desfechos de saúde ou qualidade de vida integrados, a versão documentada descreve condições regionais de calor. Ela não mede temperaturas dentro das moradias, sintomas, qualidade do sono ou impactos individuais.')
h('O que é ERA5')
p('ERA5 é a quinta geração de reanálise atmosférica do ECMWF, produzida para o Serviço de Mudanças Climáticas Copernicus. Uma reanálise combina observações históricas, como medições de estações e satélites, com um modelo físico e técnicas de assimilação de dados para estimar o estado passado da atmosfera. Oferece séries horárias globais desde 1940. [1, 2]')
p('Isso permite contextualizar as condições meteorológicas de uma região ao longo do tempo. Os valores são estimativas de reanálise, e não leituras de um termômetro instalado no ponto consultado. A resolução nativa do ERA5 é de aproximadamente 31 km; produtos distribuídos em grade regular usam 0,25°. O Open-Meteo descreve essa grade como cerca de 25 km. [1, 2, 3]')
p('O projeto acessa ERA5 pelo Open-Meteo com models=era5. Open-Meteo é a interface de acesso; ERA5 é a base meteorológica. A API pode aplicar ajustes estatísticos de altitude por padrão: o código não define elevation=nan para desativá-los. Isso não transforma a série em uma medição do microclima de Heliópolis. [3]')
p('As referências regionais ajudam a explorar intensidade, frequência e persistência do calor. Para relacioná-las à experiência dos moradores, ainda serão necessários dados locais compatíveis no tempo e no território.')

page('1 Fluxograma da sequência dos dados')
p('O fluxo tem duas entradas independentes. Os gráficos principais usam ERA5; a comparação com o entorno utiliza os dias válidos comuns à reanálise e às estações. Os valores das fontes não são fundidos.')
doc.add_picture(str(flow),width=Cm(17))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.inline_shapes[-1]._inline.docPr.set('descr','Fluxograma com duas entradas: JSON ou API ERA5 e CSVs INMET. Após validação, recorte e cálculo, os resultados aparecem no Streamlit. A comparação usa datas comuns, mantendo cada fonte separada.')
doc.add_paragraph('Figura 1  Sequência implementada de entrada, tratamento, análise e apresentação dos dados','Caption')
p('Os módulos se comunicam por chamadas de funções e objetos Python em memória. Não há banco de dados, API interna ou serviço de processamento separado.')

page('2 Coleta e tratamento da série ERA5')
h('Prioridade de carregamento')
p('Ao abrir a aplicação, iniciar_base verifica primeiro se existe um conjunto em st.session_state. Se existir, ele é reutilizado. Na primeira carga, procura dados_locais/heliopolis_open_meteo.json. Se o arquivo existir, ele é lido e validado. Somente se não existir, carregar_online consulta o Open-Meteo. Um arquivo inválido gera indisponibilidade; não há substituição automática por dados artificiais.')
p('A consulta usa urllib, da biblioteca padrão, com o endpoint https://archive-api.open-meteo.com/v1/archive. Solicita latitude −23,615 e longitude −46,590, modelo ERA5, timestamps Unix e fuso UTC. O ponto é aproximado e não corresponde a um centroide territorial oficialmente delimitado pelo projeto.')
table(['Variável solicitada','Nome interno','Unidade'],[
('Temperatura do ar a 2 m','temperatura_c','°C'),('Umidade relativa a 2 m','umidade_pct','%'),
('Vento a 10 m','vento_ms','m/s'),('Radiação solar','radiacao_wm2','W/m²'),('Precipitação','precipitacao_mm','mm')],[6.5,7.5,3])
p('A resposta JSON contém listas horárias e unidades. interpretar_api verifica os campos, converte timestamps e cria um DataFrame indexado por horário. A interface atual utiliza principalmente temperatura e umidade; vento, radiação e precipitação permanecem no processamento auxiliar.')
h('Limpeza e cobertura')
p('limpar_horarios ordena a série, remove horários duplicados mantendo a primeira ocorrência e reindexa em uma grade com todas as horas esperadas. Em 2021–2025 são 43.824 horas e 1.826 dias. Registros ausentes e valores inválidos viram NaN; não são interpolados nem substituídos por zero.')
table(['Variável','Intervalo de consistência aceito'],[('Temperatura','−90 a 65 °C'),('Umidade','0 a 100%'),('Vento','0 a 150 m/s'),('Radiação','0 a 1.600 W/m²'),('Precipitação horária','0 a 1.000 mm')],[6.5,10.5])
p('Esses limites amplos verificam consistência básica. Não são limites de saúde e não substituem controle completo de qualidade. Para cada variável, são registrados horas esperadas, válidas, ausentes ou inválidas, valores fora dos limites e percentual de cobertura.')
p('montar_conjunto reúne horários, qualidade, número de duplicatas, metadados e resposta de origem. Os metadados preservam modelo, período, coordenadas, data de coleta, fuso e URL da consulta, permitindo rastrear a proveniência.')

page('3 Seleção e tratamento das estações INMET')
h('Seleção espacial')
p('O usuário baixa os CSVs anuais e os mantém em dados_locais/inmet, com subpastas. manifesto identifica cada arquivo pelo caminho, tamanho e data de modificação. cabecalho lê o início dos arquivos para identificar código, nome, coordenadas e altitude da estação.')
p('distancia_km aplica a fórmula de Haversine entre a estação e o ponto aproximado do projeto. Usando latitudes φ, longitudes λ e diferenças em radianos, calcula:')
formula('a = sin²(Δφ / 2) + cos(φ₁) cos(φ₂) sin²(Δλ / 2)')
formula('d = 2R arcsin(√a), com R = 6.371,0088 km')
p('O raio de 30 km é uma escolha operacional do projeto, não uma norma de representatividade. Na base disponível, Mirante A701 fica a aproximadamente 13,6 km e Interlagos A771 a 15,1 km. Distância e altitude não são usadas para corrigir suas temperaturas.')
h('Leitura de cada arquivo')
p('dados.ler_inmet reconhece UTF-8 ou Latin-1, localiza o cabeçalho da tabela e interpreta separador ponto e vírgula e decimal com vírgula. Extrai data, hora UTC e temperatura do ar de bulbo seco. A integração atual usa somente essa temperatura; outras colunas do CSV não alimentam a comparação.')
p('Datas inválidas são removidas. Duplicatas dentro do mesmo CSV mantêm a primeira ocorrência. Temperaturas fora de −90 a 65 °C, inclusive o marcador −9999, são convertidas em lacunas. Um CSV sem qualquer temperatura válida é registrado como erro na auditoria.')
h('União dos anos e preservação das fontes')
p('tratar consolida os arquivos por código e localização. Variações de coordenadas de até 25 m em relação ao cadastro inicial do grupo são toleradas, com metadados originais preservados na auditoria. Mudanças maiores ficam separadas. Essa tolerância não comprova estabilidade do entorno, do instrumento ou da estação.')
p('Na união entre arquivos, valores válidos diferentes para o mesmo horário tornam o horário ausente. Valores coincidentes são consolidados. Cada série é reindexada na grade completa, mantendo lacunas e anos sem dados. Calcula-se cobertura geral e anual como horas válidas divididas pelas horas esperadas.')
p('Os horários permanecem em UTC, mas o período é delimitado pelo calendário de São Paulo. Assim, 00h, 01h e 02h UTC de 01/01/2026 podem completar 31/12/2025 local. Nenhum dia local de 2026 entra no estudo. As séries de cada estação e da ERA5 permanecem independentes.')

page('4 Cálculos diários e gráficos de temperatura')
h('Regra de cobertura diária')
p('Se n é a quantidade de registros válidos e H a quantidade de horas esperadas no dia, a cobertura é calculada pela expressão abaixo. O dia é válido com pelo menos 90% das horas de temperatura.')
formula('Cobertura diária (%) = 100 × n / H')
p('Em um dia de 24 horas, são necessários pelo menos 22 registros válidos. Um dia inválido fica com temperatura agregada e classificação de calor ausentes; ele não é interpretado como um dia ameno. Os agrupamentos usam o fuso America/Sao_Paulo.')
h('Temperatura ao longo do tempo')
p('Para cada dia válido, a média é a soma das temperaturas válidas dividida pela quantidade de registros; a máxima e a mínima são os extremos dessas amostras horárias.')
formula('Temperatura média diária = soma das temperaturas válidas / n')
p('Até 370 dias, o gráfico exibe medidas diárias. Em períodos maiores, exibe a média mensal de cada medida diária. A linha mensal de máxima é a média das máximas dos dias válidos daquele mês, não o maior valor ocorrido no mês. O mesmo vale para média e mínima. As agregações mensais não impõem um segundo limiar mínimo de dias válidos.')
p('Essa visualização permite explorar intensidade e sazonalidade. Máximas horárias podem não captar picos ocorridos entre as amostras. Cartões e episódios continuam usando medidas diárias, sem a suavização mensal do gráfico.')
h('Frequência de dias quentes')
p('Um dia válido é classificado como quente quando sua máxima é maior ou igual ao limiar L. O padrão é 30 °C, ajustável de 25 a 40 °C em passos de 0,5 °C. Esse é um critério analítico, não um limite clínico.')
formula('Frequência anual (%) = 100 × dias válidos com máxima ≥ L / dias válidos')
p('Exemplo: 70 dias acima do limiar entre 350 dias válidos representam 20%. O denominador não é automaticamente 365 ou 366. Não há um limiar adicional de cobertura anual; a frequência descreve os dias disponíveis e não estabelece tendência climática com apenas cinco anos.')
h('Calor por horário')
p('A série horária é convertida ao horário local e agrupada de 0h a 23h. Para cada hora, calcula-se a média dos registros válidos daquela hora ao longo do período. O ponto das 14h representa a média de todos os registros válidos às 14h.')
p('Esse gráfico inclui registros válidos de dias que eventualmente não atingem 90% de cobertura diária. Não é filtrado pelo limiar L. Ele descreve o perfil térmico médio ao longo da rotina, sem representar exposição pessoal medida.')

page('5 Noites umidade e persistência do calor')
h('Temperatura à noite')
p('Cada noite começa às 22h e termina às 05h59 do dia seguinte. Os registros da madrugada são associados à data de início da noite. A janela completa deve estar dentro do período consultado e ter pelo menos 90% de cobertura.')
p('Uma noite usual tem oito registros horários. Como sete correspondem a 87,5%, são necessários os oito para atingir o critério. O código também considera janelas de sete ou nove horas em transições históricas de horário de verão.')
formula('Média noturna = soma das temperaturas noturnas válidas / registros válidos')
p('Se a tabela tem mais de 370 linhas, a visualização apresenta a média mensal das médias noturnas, agrupando pela data de início da noite. Caso contrário, mostra cada noite. Noites parciais nas bordas são excluídas. A análise descreve o ambiente durante o descanso, sem medir sono ou temperatura dentro das casas.')
h('Calor e umidade')
p('Cada ponto da dispersão corresponde às médias de temperatura e umidade de um dia. O dia precisa ser válido para temperatura e ter cobertura de pelo menos 90% para umidade. A verificação é independente por variável: as horas efetivamente usadas nas duas médias podem diferir.')
p('O eixo horizontal recebe a temperatura média diária; o vertical recebe a umidade média diária. Todos os dias elegíveis aparecem, inclusive aqueles abaixo do limiar. Não há cálculo de Pearson, regressão, sensação térmica, UTCI ou índice fisiológico nessa visualização.')
h('Sequências de calor persistente')
p('episodios_quentes percorre os dias em ordem cronológica e identifica sequências de pelo menos três dias consecutivos com máxima ≥ L. Um dia abaixo do limiar, inválido ou ausente interrompe a sequência. Um episódio longo conta uma vez, e não como várias janelas sobrepostas de três dias.')
table(['Medida do episódio','Como é calculada'],[('Início e fim','Primeiro e último dia da sequência'),('Duração','Número de dias consecutivos'),('Maior temperatura','Maior máxima diária da sequência'),('Média das máximas','Média das máximas diárias da sequência'),('Borda do recorte','Indica contato com o início ou o fim da consulta')],[5.5,11.5])
p('Ao filtrar um ano, episódios que atravessam a virada podem ser cortados. A marcação de borda indica que podem continuar fora do recorte. A definição de três dias é exploratória e não corresponde automaticamente a uma onda de calor oficial.')

page('6 Cartões comparação regional e análises auxiliares')
table(['Cartão','Cálculo implementado'],[('Maior temperatura horária','Maior máxima diária entre os dias válidos'),('Dias com máxima ≥ limiar','Soma dos dias válidos classificados como quentes'),('Sequências de calor','Quantidade de episódios com três dias ou mais'),('Temperatura média noturna','Média das médias de todas as noites válidas')],[6,11])
p('No cartão noturno, cada noite válida tem o mesmo peso. Quando não há suporte suficiente, resultados ausentes são apresentados como travessão, sem confundir ausência com zero.')
h('Comparação entre ERA5 e estações')
p('contexto_regional calcula indicadores diários para cada fonte e seleciona a interseção das datas válidas. Na base atual, a comparação envolve ERA5, Mirante e Interlagos. Para cada fonte, calcula a proporção de dias com máxima ≥ L somente nesse conjunto comum.')
formula('Datas comuns = datas válidas ERA5 ∩ Mirante ∩ Interlagos')
formula('Frequência da fonte (%) = 100 × dias quentes comuns / total de datas comuns')
p('Isso reduz o efeito de lacunas sazonais: uma estação com falhas no verão não é comparada com uma série que inclui todas essas datas. A tabela informa dias comuns, dias válidos da própria fonte e distância. A frequência ERA5 nessa tabela pode diferir da calculada em toda a série porque o conjunto de datas é diferente.')
p('A correspondência atual é diária. Dentro de um dia válido, as horas ausentes podem diferir entre fontes. Não há comparação estritamente pareada de horas nessa tabela, fusão entre séries ou correção de viés para estimar o bairro.')
h('Cálculos que permanecem fora da tela atual')
p('analise.comparar_inmet continua disponível como rotina auxiliar da versão anterior. Ela compara somente pares horários simultâneos e válidos. Para Δ = temperatura ERA5 menos temperatura INMET, calcula diferença média, erro absoluto médio e raiz do erro quadrático médio:')
formula('Diferença média = média(Δ)     MAE = média(|Δ|)     RMSE = √média(Δ²)')
p('Essas métricas não alimentam a tabela regional atual. O módulo também calcula horas observadas acima do limiar, médias diárias de vento e radiação e chuva diária. A chuva exige 100% de cobertura antes de somar os valores; as demais médias auxiliares exigem 90%. Nenhuma lacuna é extrapolada.')

page('7 Responsabilidades dos arquivos principais')
h('app.py e painel.py')
p('app.py é o ponto de entrada do comando streamlit run app.py. Ele importa main de painel.py e executa a função. Não contém regras de cálculo ou coleta.')
p('painel.py coordena a aplicação. iniciar_base obtém a base; recortar_base aplica o ano selecionado; main chama indicadores_diarios, indicadores_noturnos e episodios_quentes. mostrar_exposicao desenha os cartões e gráficos de evolução e frequência; mostrar_rotina apresenta horário, noite e umidade; mostrar_persistencia exibe os episódios.')
p('A interface chama comparar_exposicao de contexto_regional.py e numero de apoio.py. Parte das agregações visuais, como médias mensais e perfil horário, ainda está em painel.py. A separação entre apresentação e cálculo é, portanto, parcial.')
h('dados.py')
p('Centraliza constantes do estudo, validação de períodos, consulta HTTP, leitura e exportação JSON, grade horária, limpeza e leitura de um CSV INMET. Não depende do Streamlit. ErroDados fornece mensagens que podem ser tratadas pela interface ou pelo terminal.')
table(['Chave do conjunto','Conteúdo'],[('horarios','DataFrame com a série horária tratada'),('qualidade','DataFrame com cobertura e inconsistências'),('duplicatas','Número de horários duplicados removidos'),('metadados','Dicionário com origem e parâmetros da coleta'),('resposta_api','Dicionário com a resposta de origem do período')],[5,12])
h('analise.py')
p('Recebe tabelas do Pandas e devolve resultados analíticos sem exigir interface. indicadores_diarios agrega dias e verifica cobertura; indicadores_noturnos organiza janelas de descanso; episodios_quentes identifica sequências; resumo_anual calcula contagens, médias e frequências. gerar_relatorio monta o texto auxiliar e comparar_inmet preserva a comparação horária legada.')
h('inmet_local.py')
p('Organiza a coleção de CSVs. manifesto identifica mudanças nos arquivos; cabecalho lê metadados; distancia_km calcula distâncias; tratar seleciona, consolida, reindexa e audita as estações, usando ler_inmet e grade_horaria de dados.py. Retorna series, catalogo, cobertura_anual e auditoria. Executado diretamente, grava resultados em dados_locais/inmet_tratado.')

page('8 Comunicação execução e persistência')
h('contexto_regional.py e apoio.py')
p('contexto_regional.py conecta as estações à análise. referencias chama tratar com cache; comparar_exposicao alinha o período, calcula indicadores com analise.py, seleciona dias comuns e apresenta a tabela. Esse arquivo reúne lógica analítica e apresentação da comparação. apoio.py contém apenas numero, responsável por casas decimais, separadores em português e ausência exibida como travessão.')
h('gerar_relatorio.py')
p('É uma entrada alternativa pelo terminal. Usa dados.py e analise.py para gerar um JSON de origem, CSVs horário, diário, noturno, episódios, resumo anual e qualidade, além de parametros.json e relatorio_inicial.txt. O argumento --arquivo permite repetir a análise offline; --inicio e --fim delimitam o período; --saida define o destino.')
p('Seus critérios são fixos em 30 °C, três dias e 90% de cobertura. O script não recebe automaticamente o limiar selecionado na sessão do dashboard e não gera a comparação regional atual. A execução pode substituir os arquivos de saída com os mesmos nomes.')
h('Reexecução e cache')
p('Ao alterar um filtro, o Streamlit executa novamente a página. st.session_state conserva o conjunto ERA5 da sessão. carregar_online usa cache de 24 horas e até seis entradas. O cache do INMET depende do manifesto com caminho, tamanho e data de modificação dos CSVs; alterações nesses atributos provocam novo tratamento.')
p('O ano altera todas as análises. O limiar altera a classificação dos dias, frequência anual, episódios e comparação regional. Ele não altera as temperaturas nem filtra os gráficos de evolução, perfil horário, noite e dispersão.')
p('A consulta online do painel não grava automaticamente o JSON em disco. Essa persistência é feita pelo script de relatório. Um JSON alterado no disco também não substitui por si só o conjunto já carregado na sessão. Uma falha na primeira carga não provoca tentativas repetidas a cada filtro na mesma sessão.')
h('Objetos que atravessam os módulos')
p('A comunicação usa funções importadas, DataFrames, dicionários e datas. Não existe uma API interna. O caminho principal é app → painel → dados e analise; a comparação segue painel → contexto_regional → inmet_local e analise. inmet_local também usa dados para ler CSVs e construir a grade. gerar_relatorio chama dados e analise sem iniciar Streamlit.')

page('9 Armazenamento testes e limites de interpretação')
table(['Arquivo ou pasta','Função'],[('dados_locais/heliopolis_open_meteo.json','Base ERA5 de entrada preferencial do painel'),('dados_locais/inmet/','CSVs originais em pastas anuais'),('dados_locais/inmet_tratado/','Exportações de séries consolidadas, cobertura e auditoria'),('Outros CSVs e TXT em dados_locais/','Resultados auxiliares gerados pelo terminal'),('requirements.txt','Dependências diretas da aplicação: Streamlit e Pandas'),('.gitignore','Exclui dados locais, ambiente virtual, caches e segredos do Git'),('.venv/ e __pycache__/','Ambiente de execução e cache de bytecode Python')],[7,10])
p('O dashboard lê o JSON e os CSVs originais INMET, recalculando resultados. Os CSVs diários exportados e a pasta inmet_tratado não são sua fonte principal. Imagens do dashboard servem para conferência visual e não alimentam a análise.')
h('Testes automatizados')
table(['Arquivo em tests/','Responsabilidade'],[('test_relatorio.py','Limpeza, cobertura, datas, noites, episódios, API e relatório'),('test_inmet_local.py','Distâncias, consolidação, conflitos e fronteiras de ano'),('test_contexto_regional.py','Comparação restrita às mesmas datas válidas'),('test_interface.py','Abertura, filtros e ausência de controles técnicos removidos')],[7,10])
p('Os testes utilizam dados sintéticos exclusivamente para validar o comportamento do código. Eles não são inseridos na base apresentada aos moradores ou pesquisadores.')
h('O que os resultados permitem afirmar')
p('A implementação permite descrever o contexto regional de calor e levantar perguntas sobre intensidade, frequência, persistência e condições noturnas. Os critérios de cobertura e a preservação de lacunas tornam explícito o suporte de dados de cada resultado.')
p('Ela não calcula efeitos clínicos, causalidade, desconforto individual, ilha de calor dentro de Heliópolis ou tendências climáticas robustas com cinco anos. O limiar de temperatura não é uma classificação nacional universal de conforto. As estações distantes não substituem a medição local e diferenças entre fontes não medem automaticamente o excesso térmico do bairro.')
p('Para cumprir integralmente o objetivo de impactos na saúde e qualidade de vida, será necessário integrar medições locais e desfechos compatíveis no tempo e no território, com definições e cobertura conhecidas. Essa etapa exige novos dados e regras analíticas, não apenas novos gráficos.')

page('10 Fontes e rastreabilidade da documentação')
p('A descrição dos módulos e cálculos foi conferida no código local vigente em 18 de outubro de 2026. As referências externas abaixo fundamentam a explicação do ERA5 e de sua disponibilização pelo Open-Meteo. Elas não substituem a verificação dos parâmetros efetivamente usados pelo projeto.')
h('Referências externas')
for text in [
    '[1] ECMWF. ECMWF Reanalysis v5 ERA5. Descrição do produto, séries horárias, origem e resolução nativa. https://www.ecmwf.int/en/forecasts/dataset/ecmwf-reanalysis-v5',
    '[2] Copernicus Climate Change Service. ERA5 hourly data on single levels from 1940 to present. Catálogo do conjunto de reanálise. https://cds.climate.copernicus.eu/datasets/reanalysis-era5-single-levels?tab=overview',
    '[3] Open-Meteo. Historical Weather API. Modelo ERA5, grade de 0,25°, variáveis, unidades, seleção de célula e ajuste estatístico por altitude. https://open-meteo.com/en/docs/historical-weather-api',
    '[4] INMET. Dados históricos anuais das estações automáticas. Origem dos CSVs disponibilizados pelo usuário. https://portal.inmet.gov.br/dadoshistoricos']:
    p(text)
p('Data de consulta das referências de ERA5 e Open-Meteo: 18 de outubro de 2026. A aplicação usa o Open-Meteo como intermediário; não faz download direto do Climate Data Store.')
h('Pontos do código usados como referência')
table(['Arquivo','Funções ou trechos centrais'],[('app.py','Inicialização de main'),('painel.py','iniciar_base, recortar_base, mostrar_exposicao, mostrar_rotina, mostrar_persistencia, main'),('dados.py','grade_horaria, limpar_horarios, interpretar_api, montar_conjunto, recortar_periodo, buscar_open_meteo, ler_inmet'),('analise.py','indicadores_diarios, indicadores_noturnos, episodios_quentes, resumo_anual, comparar_inmet'),('inmet_local.py','distancia_km, manifesto, cabecalho, tratar'),('contexto_regional.py','referencias, comparar_exposicao'),('apoio.py e gerar_relatorio.py','numero e entrada alternativa de geração de relatório')],[5.2,11.8])
p('Os nomes acima são relativos à raiz do projeto. Esta documentação registra o estado atual; mudanças futuras em filtros, cobertura, limiares, fontes ou funções devem ser refletidas no documento para preservar a reprodutibilidade.')

doc.save(OUT)
print(OUT)
